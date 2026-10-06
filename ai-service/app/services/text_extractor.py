import io
import re
from typing import Tuple


def clean_text(text: str) -> str:
    """Normalize whitespace and remove non-printable characters."""
    if not text:
        return ""
    # Replace carriage returns and excessive whitespace
    text = re.sub(r"\r\n", "\n", text)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_pdf(file_bytes: bytes) -> Tuple[bool, str, str]:
    """
    Extract readable text safely from PDF bytes.
    Returns: (success: bool, text: str, error_message: str)
    """
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        if len(reader.pages) == 0:
            return False, "", "PDF contains 0 pages"

        extracted_pages = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                extracted_pages.append(page_text.strip())

        full_text = clean_text("\n\n".join(extracted_pages))
        if not full_text:
            return False, "", "No readable text found in PDF (file may be scanned/image-only)"

        return True, full_text, ""
    except Exception as e:
        return False, "", f"Failed to extract PDF text: {str(e)}"


def extract_text_from_docx(file_bytes: bytes) -> Tuple[bool, str, str]:
    """
    Extract readable text safely from DOCX bytes.
    Returns: (success: bool, text: str, error_message: str)
    """
    try:
        import docx

        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract text inside tables
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_text.append(" | ".join(row_cells))

        combined = "\n".join(paragraphs + table_text)
        full_text = clean_text(combined)

        if not full_text:
            return False, "", "No readable text found in DOCX file"

        return True, full_text, ""
    except Exception as e:
        return False, "", f"Failed to extract DOCX text: {str(e)}"


def extract_text(file_bytes: bytes, file_type: str) -> Tuple[bool, str, str]:
    """Universal extractor handling PDF and DOCX formats."""
    file_type = file_type.lower().replace(".", "").strip()
    if file_type == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif file_type in ("docx", "doc"):
        return extract_text_from_docx(file_bytes)
    else:
        return False, "", f"Unsupported file type: {file_type}. Supported formats: PDF, DOCX"
